# -*- coding: utf-8 -*-
from django.core.management.base import BaseCommand

from datetime import datetime
from django.utils import timezone

from django.db import models

from telegram.models import Client, Order

import openpyxl


class Command(BaseCommand):
    help = 'Выгрузка статистики в excel'

    def handle(self, *args, **options):
        wb = openpyxl.Workbook()
        ws = wb.active

        ws['B1'] = 'TG ID'
        ws.column_dimensions['B'].width = 12

        ws['C1'] = 'Дата регистрации'
        ws.column_dimensions['C'].width = 16

        ws['D1'] = 'Дата первого обмена'
        ws.column_dimensions['D'].width = 18

        ws['E1'] = 'Дата последнего обмена'
        ws.column_dimensions['E'].width = 21

        ws['F1'] = 'Общее кол-во обменов'
        ws.column_dimensions['F'].width = 20

        ws['G1'] = 'Общая сумма обменов'
        ws.column_dimensions['G'].width = 20

        ws['H1'] = 'Средняя сумма обменов'
        ws.column_dimensions['H'].width = 20

        ws['I1'] = 'Макс. обмен'
        ws.column_dimensions['I'].width = 12

        ws['J1'] = 'Мин. обмен'
        ws.column_dimensions['J'].width = 12

        ws['K1'] = 'Кол-во обменов за последний месяц'
        ws.column_dimensions['K'].width = 30

        ws['L1'] = 'Сумма обменов за последний месяц'
        ws.column_dimensions['L'].width = 30

        ws['M1'] = 'Монеты'
        ws.column_dimensions['M'].width = 16

        ws['N1'] = 'Кол-во рефералов'
        ws.column_dimensions['N'].width = 16

        ws['O1'] = 'Общая сумма обменов рефералов'
        ws.column_dimensions['O'].width = 28

        ws['P1'] = 'Кол-во промокодов'
        ws.column_dimensions['P'].width = 18

        clients = Client.objects.all()

        for i, client in enumerate(clients, start=2):
            today = timezone.now()
            start_month_datetime = today.replace(day=1)

            if start_month_datetime.month == 12:
                end_month_datetime = today.replace(year=today.year + 1, month=1, day=1)

            else:
                end_month_datetime = today.replace(month=today.month + 1, day=1)

            orders = Order.objects.filter(client=client)
            month_orders = orders.filter(datetime__range=[start_month_datetime, end_month_datetime])

            register_date = client.register_date
            first_order_date = orders.aggregate(models.Min('datetime'))['datetime__min']
            last_order_date = orders.aggregate(models.Max('datetime'))['datetime__max']

            orders_count = orders.count()
            orders_sum = orders.aggregate(models.Sum('pay_value'))['pay_value__sum'] or 0.0
            orders_avg = orders.aggregate(models.Avg('pay_value'))['pay_value__avg'] or 0.0
            orders_max = orders.aggregate(models.Max('pay_value'))['pay_value__max'] or 0.0
            order_min = orders.aggregate(models.Min('pay_value'))['pay_value__min'] or 0.0

            orders_month_count = month_orders.count()
            orders_month_sum = month_orders.aggregate(models.Sum('pay_value'))['pay_value__sum'] or 0.0

            crypts = orders.values_list('crypt', flat=True).distinct()
            promocodes = len(client.get_old_promocodes() + client.get_active_promocodes()) 

            refferals = Client.objects.filter(father=client.tg_id)
            refferals_orders = Order.objects.filter(client__in=refferals)
            refferals_orders_sum = refferals_orders.aggregate(models.Sum('pay_value'))['pay_value__sum'] or 0.0

            ws[f'B{i}'] = client.tg_id
            ws[f'C{i}'] = register_date.date()
            ws[f'D{i}'] = first_order_date.date() if first_order_date is not None else '-'
            ws[f'E{i}'] = last_order_date.date() if last_order_date is not None else '-'
            ws[f'F{i}'] = orders_count
            ws[f'G{i}'] = round(orders_sum)
            ws[f'H{i}'] = round(orders_avg)
            ws[f'I{i}'] = round(orders_max)
            ws[f'J{i}'] = round(order_min)
            ws[f'K{i}'] = orders_month_count
            ws[f'L{i}'] = round(orders_month_sum)
            ws[f'M{i}'] = ','.join(crypts)
            ws[f'N{i}'] = refferals.count()
            ws[f'O{i}'] = refferals_orders_sum
            ws[f'P{i}'] = promocodes

        wb.save('report.xlsx')




